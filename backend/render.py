"""成品三件套渲染：Markdown 简历 → 专业简历 HTML → Edge 无头打印 PDF → pymupdf 转 PNG
全部本地渲染，0 token。排版对标主流专业简历（单栏、模块标题+细线、项目时间右对齐）。
模板：A=标准专业(黑字细线) B=极简留白 C=经典商务(深蓝点缀)"""
import base64
import re
import subprocess
import time
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

def _find_browser():
    """跨平台查找无头浏览器：优先 Windows Edge，其次 Linux Chromium/Chrome。"""
    candidates = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        "/usr/bin/microsoft-edge",
        "/usr/bin/microsoft-edge-stable",
        "/usr/bin/chromium",
        "/usr/bin/chromium-browser",
        "/usr/bin/google-chrome",
        "/usr/bin/google-chrome-stable",
        "/opt/google/chrome/chrome",
        "/root/.cache/ms-playwright/chromium-*/chrome-linux/chrome",
        "/home/*/.cache/ms-playwright/chromium-*/chrome-linux/chrome",
    ]
    import glob
    for p in candidates:
        if "*" in p:
            hits = glob.glob(p)
            if hits:
                return hits[0]
        elif Path(p).exists():
            return p
    return None


EDGE = _find_browser()

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_MODULE_RE = re.compile(r"^\s*#+\s+(.+?)\s*$")
_NAME_RE = re.compile(r"^\s*#\s+(.+?)\s*$")
_LIST_RE = re.compile(r"^\s*[-*•]\s+(.+?)\s*$")
_NUM_RE = re.compile(r"^\s*\d+[.、)]\s+(.+?)\s*$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_BOLD_ONLY = re.compile(r"^\*\*.+\*\*\s*$")
_META_TIME = re.compile(r"(20\d{2}\S*\s*[-–~]\s*(?:至今|现在|20\d{2}\S*))")
_META_SEP = re.compile(r"[｜|]")
_FACT_RE = re.compile(r"事实核验|【")


def _is_junk(s: str) -> bool:
    """成品里剔除：AI 自带的【事实核验】等说明性尾行"""
    return bool(_FACT_RE.search(s))


def _density(chars: int, has_photo: bool = False):
    """按内容量选密度档位，保证一页放下（A4 可排版区约 263mm）；
    有照片时整体降半档，抵消照片占用头部空间与挤窄文字区带来的行数增加"""
    if chars <= 700:
        d = dict(fs=10.5, lh=1.6, lit=1.4, hmod=4.6, head=3.5)
    elif chars <= 1000:
        d = dict(fs=10.0, lh=1.5, lit=1.2, hmod=4.2, head=3.2)
    elif chars <= 1400:
        d = dict(fs=9.5, lh=1.4, lit=0.9, hmod=3.6, head=2.9)
    elif chars <= 1900:
        d = dict(fs=9.0, lh=1.34, lit=0.75, hmod=3.2, head=2.7)
    elif chars <= 2500:
        d = dict(fs=8.5, lh=1.28, lit=0.7, hmod=3.0, head=2.5)
    else:
        d = dict(fs=8.0, lh=1.24, lit=0.6, hmod=2.8, head=2.3)
    if has_photo:
        d["fs"] = max(7.5, d["fs"] - 0.4)
        d["lh"] = max(1.15, d["lh"] - 0.05)
        d["lit"] = max(0.4, d["lit"] - 0.15)
        d["hmod"] = max(2.2, d["hmod"] - 0.3)
        d["head"] = max(2.0, d["head"] - 0.2)
    return d


def _inline(html_escaped: str) -> str:
    """**加粗** 与 *斜体* 转 <b>/<i>"""
    def bold(m):
        return f"<b>{m.group(1)}</b>"
    def ital(m):
        return f"<i>{m.group(1)}</i>"
    s = _BOLD_RE.sub(bold, html_escaped)
    s = _ITALIC_RE.sub(ital, s)
    return s


def _esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _split_meta(line: str):
    """元信息行拆成 左描述 + 右时间（如 '全栈开发者 | 2026.09 - 至今' → ('全栈开发者', '2026.09 - 至今')）"""
    m = _META_TIME.search(line)
    if m:
        return line[:m.start()].strip(), m.group(1).strip()
    return line.strip(), ""


def _parse_resume(md_text):
    """结构化解析 Markdown 简历 → (name, role, headline, blocks)
    blocks: (type, content)，type ∈ mod/proj/meta/li/num/p/table(rows)"""
    name = role = headline = ""
    blocks = []
    lines = [l.rstrip() for l in md_text.splitlines()]
    i = 0
    # 头部：姓名 / 定位行 / 一句话（在第一个模块标题或要点之前）
    head = []
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if _is_junk(s):
            i += 1
            continue
        if s.startswith(("##", "###", "-", "*", "•", "|", "1.")):
            break
        if s.startswith("#"):
            s = s[1:].strip()   # 单井号姓名（如「# 陈雨欣」）去井号，避免泄漏进成品
        head.append(s)
        i += 1
    if head:
        name = head[0]
        if len(head) > 1:
            role = head[1]
        if len(head) > 2:
            headline = "  ".join(head[2:])
    # 正文
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        blocks.append(("table", list(table_buf)))
        table_buf = []

    for raw in lines[i:]:
        line = raw.rstrip()
        s = line.strip()
        if not s:
            flush_table()
            continue
        if _is_junk(s):
            continue
        m = _MODULE_RE.match(line)
        if m:
            flush_table()
            blocks.append(("mod", m.group(1).strip()))
            continue
        m = _TABLE_ROW.match(line)
        if m:
            cells = [c.strip() for c in m.group(1).split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            table_buf.append(cells)
            continue
        flush_table()
        if _BOLD_ONLY.match(s):
            blocks.append(("proj", s))
            continue
        if _META_SEP.search(s):
            blocks.append(("meta", s))
            continue
        m = _LIST_RE.match(line)
        if m:
            blocks.append(("li", m.group(1)))
            continue
        m = _NUM_RE.match(line)
        if m:
            blocks.append(("num", m.group(1)))
            continue
        blocks.append(("p", s))
    flush_table()
    return name, role, headline, blocks


def _photo_html(photo_path):
    """照片 base64 + 原比例尺寸内联样式（不裁剪、不变形）"""
    if not photo_path or not Path(photo_path).exists():
        return ""
    try:
        b64 = base64.b64encode(Path(photo_path).read_bytes()).decode()
        mime = "image/jpeg" if str(photo_path).lower().endswith((".jpg", ".jpeg")) else "image/png"
        mm_w = mm_h = ""
        try:
            from PIL import Image
            with Image.open(photo_path) as im:
                w_px, h_px = im.size
            if w_px and h_px:
                # 原比例嵌入：宽度上限 18mm、高度上限 22mm（证件照比例），避免挤窄正文造成多页
                scale = min(22.0 / h_px, 18.0 / w_px)
                mm_w = f"width:{w_px * scale:.1f}mm;"
                mm_h = f"height:{h_px * scale:.1f}mm;"
        except Exception:
            pass
        return (f'<div class="photo"><img src="data:{mime};base64,{b64}" '
                f'style="{mm_w}{mm_h}object-fit:contain;max-width:100%"></div>')
    except Exception:
        return ""


def _css(template, d):
    """三套版式风格迥异的简历模板（一眼可辨）：
    A 现代商务 = 黑灰白·模块标题下细线·圆点符号
    B 极简杂志 = 黑白·大名字大字距·无下划线·方块符号·强留白对比
    C 经典商务 = 深蓝主色·模块标题粗线·深蓝圆点·稳重专业
    统一受密度档位压缩保证单页；项目符号由 .dot::before 按模板定制"""
    fs, lh, lit, hmod, head = d["fs"], d["lh"], d["lit"], d["hmod"], d["head"]
    base = f"""  * {{ margin:0; padding:0; box-sizing:border-box; }}
  html {{ background:#e5e7eb; }}
  @page {{ size:A4; margin:0; }}
  body {{ font-family:"Microsoft YaHei","PingFang SC","Segoe UI",sans-serif;
         width:210mm; min-height:297mm; margin:0 auto; background:#fff;
         padding:16mm 18mm 18mm; font-size:{fs}pt; line-height:{lh}; }}
  .head {{ position:relative; margin-bottom:{head}mm; padding-bottom:2.5mm;
          display:flex; justify-content:space-between; align-items:flex-start; }}
  .head-txt {{ flex:1; min-width:0; padding-right:5mm; }}
  .role {{ font-size:10.5pt; color:#374151; margin-top:1.2mm; font-weight:600; }}
  .headline {{ font-size:9.5pt; color:#6b7280; margin-top:0.8mm; }}
  .photo {{ flex:0 0 auto; }}
  .photo img {{ object-fit:contain; border-radius:1.5mm; }}
  .hmod {{ font-size:12.5pt; font-weight:700; color:#111827; letter-spacing:1px;
          margin:{hmod}mm 0 1.8mm; padding-bottom:1mm; }}
  .proj {{ font-size:11pt; font-weight:700; color:#111827; margin:2.6mm 0 0.8mm; }}
  .meta {{ display:flex; justify-content:space-between; font-size:9pt; color:#6b7280;
          margin-bottom:1.2mm; }}
  .mt {{ color:#6b7280; }}
  .lit {{ display:flex; gap:2.4mm; margin-bottom:{lit}mm; font-size:10pt;
          align-items:flex-start; }}
  .dot {{ flex:0 0 auto; width:3.2mm; text-align:center; line-height:1; margin-top:0.5mm; }}
  b {{ color:#111827; }}
  table {{ width:100%; border-collapse:collapse; margin:1.5mm 0; }}
  td {{ border:0.5pt solid #d1d5db; padding:1.2mm 2mm; font-size:9pt; }}
  .par {{ margin-bottom:1.5mm; font-size:10pt; }}
  @media print {{ html, body {{ background:#fff !important; }} }}
  body {{ -webkit-print-color-adjust:exact; print-color-adjust:exact; }}"""

    if template == "B":
        # 极简杂志：28pt 超大字距名字、无下划线、模块标题细字重大间距、方块符号
        # 通过更大标题留白 + 更紧正文制造"杂志式"强对比，同时保证单页
        css = base + f"""
  .hname {{ font-size:28pt; font-weight:200; letter-spacing:10px; color:#111827;
          line-height:1.15; }}
  .role {{ font-size:10pt; font-weight:600; color:#111827; letter-spacing:2px; }}
  .headline {{ font-size:9pt; color:#9ca3af; letter-spacing:0.5px; }}
  .head {{ border-bottom:0; margin-bottom:{head + 2.5:.1f}mm; }}
  .hmod {{ font-size:14pt; font-weight:600; letter-spacing:5px; color:#111827;
          margin:{hmod + 1.6:.1f}mm 0 2.6mm; padding-bottom:0; border-bottom:0; }}
  .dot {{ width:3.8mm; }}
  .dot::before {{ content:""; display:block; width:1.4mm; height:1.4mm;
          background:#111827; border-radius:0.2mm; margin:1.4mm auto 0; }}
  .proj {{ font-size:11.5pt; font-weight:700; color:#111827; margin-top:3.2mm; }}
  .lit {{ gap:2.8mm; margin-bottom:{max(0.35, lit - 0.25):.2f}mm; font-size:{max(8.5, fs - 0.3):.1f}pt; }}
  .meta {{ color:#9ca3af; }}
  table td {{ border:0.3pt solid #e5e7eb; }}"""
        return css

    if template == "C":
        # 经典商务：深蓝主色 + 头部浅蓝底块 + 模块粗蓝线 + 深蓝圆点
        css = base + f"""
  .hname {{ font-size:22pt; font-weight:800; letter-spacing:1px; color:#1f4e79; }}
  .role {{ color:#1f4e79; font-weight:700; }}
  .headline {{ color:#5b7a94; }}
  .head {{ background:#eef3f8; padding:6mm 7mm 4.5mm; margin-bottom:5.5mm;
          border-bottom:1.2pt solid #1f4e79; }}
  .head-txt {{ padding-right:6mm; }}
  .photo img {{ border-radius:1mm; }}
  .hmod {{ font-size:12.5pt; font-weight:700; letter-spacing:1px; color:#1f4e79;
          border-bottom:1pt solid #1f4e79; padding-bottom:1.2mm; }}
  .dot::before {{ content:"•"; font-weight:700; color:#1f4e79; }}
  .proj {{ color:#1f4e79; }}
  .meta {{ color:#4b6e8a; }}
  .mt {{ color:#4b6e8a; }}
  b {{ color:#1f4e79; }}
  table td {{ border:0.6pt solid #9fb6cc; }}
  .lit {{ margin-bottom:{max(0.35, lit - 0.15):.2f}mm; font-size:{max(8.5, fs - 0.2):.1f}pt; }}"""
        return css

    # A 现代商务：黑灰白 · 模块标题下细灰线 · 深灰圆点
    css = base + """
  .hname { font-size:21pt; font-weight:800; letter-spacing:2px; color:#111827; }
  .head { border-bottom:0.5pt solid #d1d5db; }
  .hmod { font-size:12.5pt; font-weight:700; letter-spacing:1.5px; color:#111827;
          border-bottom:0.5pt solid #9ca3af; padding-bottom:1mm; }
  .dot::before { content:"•"; font-weight:700; color:#374151; }
  .proj { color:#111827; }
  .meta { color:#6b7280; }"""
    return css


def md_to_html(md_text: str, photo_path=None, name_hint="", template="A") -> str:
    """Markdown 简历 → 专业简历 HTML（照片原比例内嵌，可直接打印）"""
    name, role, headline, blocks = _parse_resume(md_text)
    if not name and name_hint:
        name = name_hint
    photo_html = _photo_html(photo_path)
    # 头部
    head_html = '<div class="head"><div class="head-txt">'
    if name:
        head_html += f'<div class="hname">{_esc(name)}</div>'
    if role:
        head_html += f'<div class="role">{_esc(role)}</div>'
    if headline:
        head_html += f'<div class="headline">{_esc(headline)}</div>'
    head_html += "</div>" + photo_html + "</div>"
    # 正文（content 为纯 markdown 文本，渲染时统一转 HTML）
    body = []
    for typ, content in blocks:
        if typ == "mod":
            body.append(f'<div class="hmod">{_inline(_esc(content))}</div>')
        elif typ == "proj":
            body.append(f'<div class="proj">{_inline(_esc(content))}</div>')
        elif typ == "meta":
            left, right = _split_meta(content)
            inner = f'<span>{_inline(_esc(left))}</span>'
            if right:
                inner += f'<span class="mt">{_esc(right)}</span>'
            body.append(f'<div class="meta">{inner}</div>')
        elif typ in ("li", "num"):
            body.append(f'<div class="lit"><span class="dot"></span><span>{_inline(_esc(content))}</span></div>')
        elif typ == "table":
            rows = "".join(
                "<tr>" + "".join(f"<td>{_inline(_esc(c.strip()))}</td>" for c in r) + "</tr>"
                for r in content)
            body.append(f"<table>{rows}</table>")
        else:
            body.append(f'<div class="par">{_inline(_esc(content))}</div>')
    title = f"{name} - 岗位定制简历" if name else "岗位定制简历"
    tpl = (template or "A").upper()
    if tpl not in ("A", "B", "C"):
        tpl = "A"
    d = _density(len(md_text), has_photo=photo_html is not None)
    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{_esc(title)}</title>
<style>
{_css(tpl, d)}
</style></head><body>
{head_html}
{''.join(body)}
</body></html>"""


def html_to_pdf(html: str, out_pdf: Path) -> Path:
    """无头浏览器打印 HTML → PDF（A4）。跨平台：Edge → Chromium/Chrome → Playwright。"""
    import tempfile
    tmp_html = Path(tempfile.gettempdir()) / f"_rz_resume_{int(time.time())}.html"
    tmp_html.write_text(html, encoding="utf-8")

    def _wait_pdf():
        for _ in range(30):
            if out_pdf.exists() and out_pdf.stat().st_size > 1024:
                return True
            time.sleep(0.5)
        return False

    def _via_browser(browser):
        args = [
            browser, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
            "--print-background",
            "--print-to-pdf=" + str(out_pdf),
            tmp_html.as_uri(),
        ]
        subprocess.run(args, capture_output=True, timeout=90)
        return _wait_pdf()

    browser = EDGE or _find_browser()
    if browser and _via_browser(browser):
        tmp_html.unlink(missing_ok=True)
        return out_pdf

    # 兜底：Playwright 无头 Chromium
    try:
        from playwright.sync_api import sync_playwright
        with sync_playwright() as p:
            b = p.chromium.launch(args=["--no-sandbox", "--disable-gpu"])
            page = b.new_page()
            page.goto(tmp_html.as_uri())
            page.pdf(path=str(out_pdf), print_background=True, format="A4",
                     margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
            b.close()
        ok = _wait_pdf()
    except Exception:
        ok = False

    tmp_html.unlink(missing_ok=True)
    if not ok:
        raise RuntimeError("未找到可用的无头浏览器（Edge/Chromium/Playwright），无法生成 PDF")
    return out_pdf


def pdf_to_png(pdf_path: Path, out_dir: Path, dpi=150) -> list:
    """PDF 每页 → PNG，返回 PNG 路径列表"""
    import fitz
    out_dir.mkdir(parents=True, exist_ok=True)
    pngs = []
    doc = fitz.open(str(pdf_path))
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=dpi)
        png = out_dir / f"{pdf_path.stem}_p{i + 1}.png"
        pix.save(str(png))
        pngs.append(png)
    doc.close()
    return pngs


# ---------------- Word（与三件套共用逻辑，含照片） ----------------
def _add_runs(paragraph, text):
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def _para_bottom_border(para, sz=8, color="374151"):
    """段落底部加细线（模块标题分隔线）"""
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def markdown_to_word(md_text, out_path, photo_path=None):
    """Markdown 简历 → 专业 Word（顶部双栏：左姓名/定位，右照片；模块标题+细线；项目时间右对齐；单页密度自适应）"""
    doc = Document()
    for sec in doc.sections:
        sec.top_margin = Cm(1.5)
        sec.bottom_margin = Cm(1.5)
        sec.left_margin = Cm(1.8)
        sec.right_margin = Cm(1.8)
    d = _density(len(md_text), has_photo=bool(photo_path and Path(photo_path).exists()))
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(d["fs"])
    style.paragraph_format.line_spacing = d["lh"]
    name, role, headline, blocks = _parse_resume(md_text)

    # 顶部：左信息 + 右照片
    tbl = doc.add_table(rows=1, cols=2)
    tbl.autofit = True
    left_cell, right_cell = tbl.cell(0, 0), tbl.cell(0, 1)
    if name:
        p = left_cell.paragraphs[0]
        r = p.add_run(name)
        r.bold = True
        r.font.size = Pt(19)
        r.font.name = "微软雅黑"
        p.paragraph_format.space_after = Pt(2)
    if role:
        p2 = left_cell.add_paragraph()
        r2 = p2.add_run(role)
        r2.bold = True
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        p2.paragraph_format.space_after = Pt(1)
    if headline:
        p3 = left_cell.add_paragraph()
        r3 = p3.add_run(headline)
        r3.font.size = Pt(9.5)
        r3.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
        p3.paragraph_format.space_after = Pt(0)
    if photo_path and Path(photo_path).exists():
        try:
            pp = right_cell.paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pp.add_run().add_picture(str(photo_path), width=Inches(1.1))
        except Exception:
            pass

    # 正文
    for typ, content in blocks:
        if typ == "mod":
            hp = doc.add_paragraph()
            hr = hp.add_run(content)
            hr.bold = True
            hr.font.size = Pt(12.5)
            hr.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            hp.paragraph_format.space_before = Pt(max(5, d["hmod"] * 2.4))
            hp.paragraph_format.space_after = Pt(2)
            _para_bottom_border(hp)
        elif typ == "proj":
            pp = doc.add_paragraph()
            _add_runs(pp, content)
            for r in pp.runs:
                r.bold = True
                r.font.size = Pt(11)
            pp.paragraph_format.space_before = Pt(4)
            pp.paragraph_format.space_after = Pt(1)
        elif typ == "meta":
            left, right = _split_meta(content)
            mp = doc.add_paragraph()
            mp.paragraph_format.space_after = Pt(1)
            if right:
                mp.paragraph_format.tab_stops.add_tab_stop(
                    Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
            _add_runs(mp, left)
            for r in mp.runs:
                r.font.size = Pt(9)
                r.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
            if right:
                _add_runs(mp, "\t" + right)
                for r in mp.runs:
                    r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x6b, 0x72, 0x80)
        elif typ in ("li", "num"):
            lp = doc.add_paragraph()
            lp.paragraph_format.left_indent = Cm(0.5)
            lp.paragraph_format.space_after = Pt(1.5)
            _add_runs(lp, "• " + content)
        elif typ == "table":
            rows = content
            if rows:
                wt = doc.add_table(rows=len(rows), cols=len(rows[0]))
                wt.style = "Table Grid"
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        wt.cell(ri, ci).text = cell.strip()
        else:
            pp = doc.add_paragraph()
            _add_runs(pp, content)
            pp.paragraph_format.space_after = Pt(1.5)
    doc.save(out_path)
    return out_path


def extract_docx_photo(docx_path, out_dir):
    """从 docx 简历提取第一张照片到 out_dir，返回路径；无照片返回 None"""
    try:
        doc = Document(docx_path)
        for shape in doc.inline_shapes:
            t = getattr(shape, "type", None)
            if getattr(t, "name", "") != "PICTURE":
                continue
            inline = getattr(shape, "_inline", None)
            if inline is None:
                continue
            blips = inline.findall(".//" + qn("a:blip"))
            if not blips:
                continue
            embed = blips[0].get(qn("r:embed"))
            if not embed:
                continue
            part = doc.part.related_parts[embed]
            blob = part.blob
            ext = getattr(part, "ext", None)
            if not ext:
                ctype = getattr(part, "content_type", "") or ""
                ext = ctype.split("/")[-1] if "/" in ctype else "png"
            ext = ext or "png"
            target = Path(out_dir) / f"resume_photo.{ext}"
            target.write_bytes(blob)
            return target
    except Exception:
        pass
    return None


def extract_pdf_photo(pdf_path, out_dir):
    """从 PDF 简历第一页提取图片（通常为证件照/插图），返回路径；无图片返回 None"""
    try:
        import fitz
        doc = fitz.open(str(pdf_path))
        page = doc[0]
        for img in page.get_images(full=True):
            xref = img[0]
            pix = fitz.Pixmap(doc, xref)
            if pix.n - pix.alpha > 3:  # CMYK 转 RGB
                pix = fitz.Pixmap(fitz.csRGB, pix)
            ext = "png" if pix.alpha else "jpeg"
            target = Path(out_dir) / f"resume_photo.{ext}"
            pix.save(str(target))
            pix = None
            doc.close()
            return target
        doc.close()
    except Exception:
        pass
    return None
