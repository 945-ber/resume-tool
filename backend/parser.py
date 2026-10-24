"""简历解析模块：支持 .docx / .pdf / 图片 / .txt"""
from pathlib import Path


def parse_docx(path):
    from docx import Document
    doc = Document(path)
    parts = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_pdf(path):
    import pymupdf  # PyMuPDF
    doc = pymupdf.open(path)
    parts = []
    total_chars = 0
    for page in doc:
        text = page.get_text("text").strip()
        total_chars += len(text)
        if text:
            parts.append(text)
    return "\n".join(parts), total_chars


def parse_file(path):
    ext = Path(path).suffix.lower()
    if ext == ".docx":
        text = parse_docx(path)
        return {"format": "docx", "text": text, "needs_ocr": False}
    if ext == ".pdf":
        text, n = parse_pdf(path)
        if n < 20:
            return {"format": "pdf_scan", "text": "", "needs_ocr": True}
        return {"format": "pdf", "text": text, "needs_ocr": False}
    if ext in (".png", ".jpg", ".jpeg", ".bmp", ".webp"):
        return {"format": "image", "text": "", "needs_ocr": True}
    if ext == ".txt":
        with open(path, encoding="utf-8", errors="ignore") as f:
            return {"format": "txt", "text": f.read(), "needs_ocr": False}
    return {"format": "unsupported", "text": "", "needs_ocr": False}
