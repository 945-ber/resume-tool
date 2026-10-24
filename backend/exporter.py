"""导出：Markdown 简历 → 干净排版 Word 文档（内联标记转真实格式，可插入原简历照片）"""
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")
_ITALIC_RE = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")


def _add_runs(paragraph, text):
    """把含 **加粗** / *斜体* 的文本按格式拆分写入段落 run，不留星号"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            _add_plain_or_italic(paragraph, text[pos:m.start()])
        paragraph.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        _add_plain_or_italic(paragraph, text[pos:])


def _add_plain_or_italic(paragraph, text):
    pos = 0
    for m in _ITALIC_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        paragraph.add_run(m.group(1)).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def markdown_to_word(md_text, out_path, photo_path=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "微软雅黑"
    style.font.size = Pt(10.5)

    # 照片（个人信息处，居中小尺寸）
    if photo_path:
        try:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(photo_path), width=Inches(1.5))
        except Exception:
            pass  # 照片插入失败不影响文档

    for raw in md_text.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\s*[-*•]\s+", line):
            para = doc.add_paragraph(style="List Bullet")
            _add_runs(para, re.sub(r"^\s*[-*•]\s+", "", line))
        elif re.match(r"^\s*\d+[.、)]\s+", line):
            para = doc.add_paragraph(style="List Number")
            _add_runs(para, re.sub(r"^\s*\d+[.、)]\s+", "", line))
        else:
            para = doc.add_paragraph()
            _add_runs(para, line)
    doc.save(out_path)
    return out_path


def extract_docx_photo(docx_path, out_dir):
    """从 docx 简历中提取第一张照片到 out_dir，返回路径；无照片返回 None"""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        doc = Document(docx_path)
        for shape in doc.inline_shapes:
            t = getattr(shape, "type", None)
            if getattr(t, "name", "") != "PICTURE":
                continue
            inline = getattr(shape, "_inline", None)
            if inline is None:
                continue
            blips = inline.findall(".//" + qn("a:blip"))
            if not blips:
                continue
            embed = blips[0].get(qn("r:embed"))
            if not embed:
                continue
            part = doc.part.related_parts[embed]
            blob = part.blob
            ext = getattr(part, "ext", None)
            if not ext:
                ctype = getattr(part, "content_type", "") or ""
                ext = ctype.split("/")[-1] if "/" in ctype else "png"
            ext = ext or "png"
            target = out_dir / f"resume_photo.{ext}"
            target.write_bytes(blob)
            return target
    except Exception:
        pass
    return None
